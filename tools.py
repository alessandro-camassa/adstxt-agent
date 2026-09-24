"""Tools for the adstxt agent.

Intake: inspect_excel and stage_import are LangChain tools the intake agent can
call. commit_import is plain code: the model never calls it and never writes to
the domains table.

Checks: check_domains and check_domain fetch ads.txt with curl, match our
OneTag line in code and write to the results and manual_check tables.
"""

import json
import re
import sqlite3
import subprocess
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

from langchain_core.tools import tool
from openpyxl import load_workbook

ROOT = Path(__file__).parent
DB_PATH = ROOT / "data" / "adstxt.db"

PREVIEW_ROWS = 5
DOMAIN_RE = re.compile(r"^(?=.{1,253}$)([a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}$")
TOTAL_RE = re.compile(r"^\s*(grand\s+)?total[ei]?\b", re.IGNORECASE)
MIRROR_SUFFIX = ".translate.goog"


# ---------- helpers ----------

def _cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def _is_blank_row(row) -> bool:
    return all(_cell_text(v) == "" for v in row)


def _is_total_row(row) -> bool:
    """A row labelled "Total", or an unlabelled one with only numbers and no text
    (ad server exports often put the grand total there, above or below the data).
    commit_import still checks that it really equals the sum of the other rows."""
    if any(isinstance(v, str) and TOTAL_RE.match(v) for v in row):
        return True
    filled = [v for v in row if _cell_text(v)]
    return bool(filled) and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in filled)


def _normalize_domain(value) -> str:
    return _cell_text(value).lower()


def _mirror_origin(domain: str) -> str | None:
    """www-example-com.translate.goog -> www.example.com"""
    if not domain.endswith(MIRROR_SUFFIX):
        return None
    encoded = domain[: -len(MIRROR_SUFFIX)]
    return encoded.replace("--", "\0").replace("-", ".").replace("\0", "-")


def _parse_number(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _cell_text(value).replace(",", "").replace(" ", "").replace("_", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _guess_header_row(rows) -> int | None:
    """First of the top 20 rows with the most filled cells, mostly text (skips titles above the table)."""
    top = rows[:20]
    if not top:
        return None
    filled = {n: [v for v in r if _cell_text(v)] for n, r in top}
    widest = max(len(v) for v in filled.values())
    for n, values in filled.items():
        if len(values) == widest and sum(isinstance(v, str) for v in values) * 2 >= len(values):
            return n
    return top[0][0]


def _read_sheet(path: str, sheet: str | None = None, header_row: int | None = None):
    """Return (sheet_name, header_row_number, headers, data_rows).

    If header_row is not given it is guessed. data_rows is a list of
    (excel_row_number, values) for every non-blank row after the header.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = [(i, list(values)) for i, values in enumerate(ws.iter_rows(values_only=True), start=1)
                if not _is_blank_row(values)]
        title = ws.title
    finally:
        wb.close()

    if header_row is None:
        header_row = _guess_header_row(rows)
    headers = next(([_cell_text(v) for v in r] for n, r in rows if n == header_row), [])
    return title, header_row, headers, [(n, r) for n, r in rows if header_row and n > header_row]


def _column_index(headers: list[str], name: str | None) -> int | None:
    if not name:
        return None
    wanted = name.strip().lower()
    for i, h in enumerate(headers):
        if h.lower() == wanted:
            return i
    raise ValueError(f"Column {name!r} not found. Headers are: {headers}")


def _value(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


def _guess_domain_column(headers, rows) -> int | None:
    best, best_score = None, 0
    for i in range(len(headers)):
        values = [_normalize_domain(_value(r, i)) for _, r in rows[:200]]
        score = sum(1 for v in values if DOMAIN_RE.match(v))
        if score > best_score:
            best, best_score = i, score
    return best


def _connect() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS imports (
            import_id     INTEGER PRIMARY KEY AUTOINCREMENT,
            path          TEXT NOT NULL,
            sheet         TEXT NOT NULL,
            mapping       TEXT NOT NULL,
            rows_read     INTEGER NOT NULL,
            file_total    REAL,
            total_row     INTEGER,
            staged_at     TEXT NOT NULL,
            committed_at  TEXT
        );
        CREATE TABLE IF NOT EXISTS staging (
            import_id     INTEGER NOT NULL REFERENCES imports(import_id),
            row_number    INTEGER NOT NULL,
            raw_domain    TEXT,
            domain        TEXT,
            bid_requests  REAL,
            publisher     TEXT,
            status        TEXT NOT NULL CHECK (status IN ('accepted', 'rejected')),
            reason        TEXT,
            PRIMARY KEY (import_id, row_number)
        );
        CREATE TABLE IF NOT EXISTS domains (
            domain        TEXT PRIMARY KEY,
            bid_requests  REAL NOT NULL,
            publisher     TEXT,
            import_id     INTEGER NOT NULL REFERENCES imports(import_id),
            imported_at   TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS results (
            domain            TEXT PRIMARY KEY,
            checked_domain    TEXT NOT NULL,
            used_main_domain  INTEGER NOT NULL DEFAULT 0,
            status            TEXT NOT NULL CHECK (status IN ('authorised', 'ob_only',
                                'other_onetag_id', 'no_onetag_line', 'no_ads_txt',
                                'blocked', 'mirror', 'error')),
            http_code         INTEGER,
            final_url         TEXT,
            evidence          TEXT,
            relationship      TEXT,
            detail            TEXT,
            bid_requests      REAL,
            checked_at        TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS manual_check (
            domain        TEXT PRIMARY KEY,
            bid_requests  REAL,
            url           TEXT,
            http_code     INTEGER,
            reason        TEXT,
            added_at      TEXT NOT NULL
        );
        """
    )
    return conn


# ---------- tools the agent can call ----------

@tool
def inspect_excel(path: str) -> str:
    """Profile an Excel file without returning every row.

    Returns sheet names, headers, row count, the first and last few rows, and
    anything odd: total rows, blank or non-domain values, duplicate domains and
    translate.goog mirrors. Call this first, before stage_import.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    sheets = []
    for name in sheet_names:
        title, header_row, headers, rows = _read_sheet(path, name)
        total_rows = [n for n, r in rows if _is_total_row(r)]
        data = [(n, r) for n, r in rows if n not in total_rows]
        dom_idx = _guess_domain_column(headers, data)

        odd = {"total_rows": [], "blank_cells_per_column": {}}
        for n, r in rows:
            if n in total_rows:
                odd["total_rows"].append({"row": n, "values": [_cell_text(v) for v in r]})
        for i, h in enumerate(headers):
            blanks = sum(1 for _, r in data if _cell_text(_value(r, i)) == "")
            if blanks:
                odd["blank_cells_per_column"][h or f"column_{i + 1}"] = blanks

        if dom_idx is not None:
            domains = [(n, _normalize_domain(_value(r, dom_idx))) for n, r in data]
            non_domain = [
                {"row": n, "value": d} for n, d in domains
                if d and not DOMAIN_RE.match(d)
            ]
            counts = Counter(d for _, d in domains if d)
            mirrors = [
                {"row": n, "value": d, "origin": _mirror_origin(d)}
                for n, d in domains if d.endswith(MIRROR_SUFFIX)
            ]
            odd["likely_domain_column"] = headers[dom_idx]
            odd["non_domain_values"] = {"count": len(non_domain), "examples": non_domain[:10]}
            odd["duplicate_domains"] = {
                "count": sum(c - 1 for c in counts.values() if c > 1),
                "examples": [{"domain": d, "times": c} for d, c in counts.most_common(10) if c > 1],
            }
            odd["translate_goog_mirrors"] = {"count": len(mirrors), "examples": mirrors[:10]}

        sheets.append({
            "sheet": title,
            "header_row": header_row,
            "headers": headers,
            "data_row_count": len(data),
            "first_rows": [[_cell_text(v) for v in r] for _, r in data[:PREVIEW_ROWS]],
            "last_rows": [[_cell_text(v) for v in r] for _, r in data[-PREVIEW_ROWS:]],
            "odd": odd,
        })

    return json.dumps({"path": path, "sheet_names": sheet_names, "sheets": sheets}, default=str)


@tool
def stage_import(path: str, mapping: dict) -> str:
    """Extract every row of the file into the staging table.

    mapping must have "domain" and "bid_requests" (header names of those
    columns) and may have "publisher", "sheet" (defaults to the first sheet)
    and "header_row" (the Excel row number of the headers, as reported by
    inspect_excel). Each row is marked accepted or rejected with a reason. The
    file's own total row is kept aside for reconciliation. Nothing is written
    to the domains table here.
    """
    sheet, header_row, headers, rows = _read_sheet(
        path, mapping.get("sheet"), int(mapping["header_row"]) if mapping.get("header_row") else None)
    dom_idx = _column_index(headers, mapping.get("domain"))
    bid_idx = _column_index(headers, mapping.get("bid_requests"))
    pub_idx = _column_index(headers, mapping.get("publisher"))
    if dom_idx is None or bid_idx is None:
        raise ValueError("mapping needs both 'domain' and 'bid_requests'")

    total_rows = [(n, r) for n, r in rows if _is_total_row(r)]
    file_total, total_row = None, None
    if total_rows:
        total_row, r = total_rows[-1]
        file_total = _parse_number(_value(r, bid_idx))
    data = [(n, r) for n, r in rows if not _is_total_row(r)]

    staged, seen = [], {}
    for n, r in data:
        raw = _cell_text(_value(r, dom_idx))
        domain = _normalize_domain(raw)
        bids = _parse_number(_value(r, bid_idx))
        publisher = _cell_text(_value(r, pub_idx)) or None

        reason = None
        if not domain:
            reason = "blank domain"
        elif domain.endswith(MIRROR_SUFFIX):
            reason = f"translate.goog mirror of {_mirror_origin(domain)}"
        elif not DOMAIN_RE.match(domain):
            reason = "not a valid domain"
        elif domain in seen:
            reason = f"duplicate of row {seen[domain]}"
        elif bids is None:
            reason = "bid requests not a number"
        if domain and domain not in seen:
            seen[domain] = n

        staged.append((n, raw, domain or None, bids, publisher,
                       "rejected" if reason else "accepted", reason))

    conn = _connect()
    with conn:
        cur = conn.execute(
            "INSERT INTO imports (path, sheet, mapping, rows_read, file_total, total_row, staged_at)"
            " VALUES (?, ?, ?, ?, ?, ?, ?)",
            (str(path), sheet, json.dumps(mapping), len(data), file_total, total_row,
             datetime.now(timezone.utc).isoformat()),
        )
        import_id = cur.lastrowid
        conn.executemany(
            "INSERT INTO staging (import_id, row_number, raw_domain, domain, bid_requests,"
            " publisher, status, reason) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            [(import_id, *s) for s in staged],
        )
    conn.close()

    accepted = sum(1 for s in staged if s[5] == "accepted")
    reasons = Counter(s[6].split(" of ")[0] for s in staged if s[6])
    return json.dumps({
        "import_id": import_id,
        "sheet": sheet,
        "rows_read": len(data),
        "accepted": accepted,
        "rejected": len(staged) - accepted,
        "rejection_reasons": dict(reasons),
        "file_total_row": total_row,
        "file_total_bid_requests": file_total,
    })


# ---------- plain code, never called by the model ----------

def commit_import(import_id: int | None = None) -> dict:
    """Move accepted staged rows into the domains table, only if the numbers reconcile.

    Checks that rows read == accepted + rejected, and that bid requests across
    accepted and rejected rows together equal the file's own total row.
    Defaults to the most recent staged import.
    """
    conn = _connect()
    try:
        if import_id is None:
            row = conn.execute("SELECT MAX(import_id) FROM imports").fetchone()
            import_id = row[0]
        if import_id is None:
            return {"committed": False, "reasons": ["Nothing has been staged yet."]}

        imp = conn.execute("SELECT * FROM imports WHERE import_id = ?", (import_id,)).fetchone()
        if imp["committed_at"]:
            return {"import_id": import_id, "committed": False,
                    "reasons": [f"Import {import_id} was already committed."]}

        counts = {r["status"]: r for r in conn.execute(
            "SELECT status, COUNT(*) AS n, COALESCE(SUM(bid_requests), 0) AS bids,"
            " SUM(bid_requests IS NULL) AS missing"
            " FROM staging WHERE import_id = ? GROUP BY status", (import_id,))}
        accepted = counts["accepted"]["n"] if "accepted" in counts else 0
        rejected = counts["rejected"]["n"] if "rejected" in counts else 0
        staged_bids = sum(c["bids"] for c in counts.values())
        missing_bids = sum(c["missing"] for c in counts.values())
        file_total = imp["file_total"]

        checks = {
            "rows_read": imp["rows_read"],
            "rows_accepted": accepted,
            "rows_rejected": rejected,
            "rows_match": imp["rows_read"] == accepted + rejected,
            "bid_requests_staged": staged_bids,
            "bid_requests_file_total": file_total,
            "rows_without_bid_number": missing_bids,
            "bids_match": file_total is not None and abs(staged_bids - file_total) < 0.5,
        }

        reasons = []
        if not checks["rows_match"]:
            reasons.append(
                f"Rows don't add up: {imp['rows_read']} read but "
                f"{accepted} accepted + {rejected} rejected = {accepted + rejected}.")
        if file_total is None:
            reasons.append("The file has no total row (or its bid requests cell is not a number), "
                           "so there is nothing to reconcile against.")
        elif not checks["bids_match"]:
            msg = (f"Bid requests don't add up: accepted + rejected rows sum to "
                   f"{staged_bids:,.0f} but the file's total row says {file_total:,.0f} "
                   f"(difference {staged_bids - file_total:+,.0f}).")
            if missing_bids:
                msg += f" {missing_bids} row(s) have no readable bid requests number."
            reasons.append(msg)

        if reasons:
            return {"import_id": import_id, "committed": False, "reasons": reasons, "checks": checks}

        now = datetime.now(timezone.utc).isoformat()
        with conn:
            conn.execute(
                "INSERT INTO domains (domain, bid_requests, publisher, import_id, imported_at)"
                " SELECT domain, bid_requests, publisher, import_id, ? FROM staging"
                " WHERE import_id = ? AND status = 'accepted'"
                " ON CONFLICT(domain) DO UPDATE SET bid_requests = excluded.bid_requests,"
                " publisher = excluded.publisher, import_id = excluded.import_id,"
                " imported_at = excluded.imported_at",
                (now, import_id),
            )
            conn.execute("UPDATE imports SET committed_at = ? WHERE import_id = ?", (now, import_id))
        return {"import_id": import_id, "committed": True, "reasons": [], "checks": checks}
    finally:
        conn.close()


def import_summary(import_id: int) -> dict:
    """Mapping, counts and rejected rows for one import, for display."""
    conn = _connect()
    try:
        imp = conn.execute("SELECT * FROM imports WHERE import_id = ?", (import_id,)).fetchone()
        rejected = [dict(r) for r in conn.execute(
            "SELECT row_number, raw_domain, bid_requests, publisher, reason FROM staging"
            " WHERE import_id = ? AND status = 'rejected' ORDER BY row_number", (import_id,))]
        accepted = conn.execute(
            "SELECT COUNT(*) FROM staging WHERE import_id = ? AND status = 'accepted'",
            (import_id,)).fetchone()[0]
        return {
            "import_id": import_id,
            "sheet": imp["sheet"],
            "mapping": json.loads(imp["mapping"]),
            "rows_read": imp["rows_read"],
            "accepted": accepted,
            "rejected_rows": rejected,
        }
    finally:
        conn.close()


def latest_import_id() -> int | None:
    conn = _connect()
    try:
        return conn.execute("SELECT MAX(import_id) FROM imports").fetchone()[0]
    finally:
        conn.close()


# ---------- ads.txt checks ----------

ONETAG_DOMAIN = "onetag.com"
ONETAG_ID = "65e2f0d9f4ee117"
ONETAG_OB_ID = f"{ONETAG_ID}-ob"
CONCURRENCY = 20
TIMEOUT_SECONDS = 10
BLOCK_RETRY_WAIT = 3
USER_AGENT = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36")
CURL_MARKER = "\n__curl__"
BOT_CHECK_MARKERS = ("just a moment", "cf-browser-verification", "challenge-platform", "captcha",
                     "attention required", "access denied", "are you a robot", "are you human",
                     "checking your browser", "ddos protection", "bot detection", "request unsuccessful",
                     "awswaf", "_incapsula_", "perimeterx", "datadome")
# Suffixes under which the registrable domain has three labels (not exhaustive).
MULTI_PART_SUFFIXES = {
    "co.uk", "org.uk", "ac.uk", "gov.uk", "com.au", "net.au", "org.au", "co.nz", "co.jp",
    "co.za", "com.br", "com.mx", "com.ar", "com.tr", "co.in", "co.kr", "com.sg", "com.hk",
    "com.cn", "net.cn", "org.cn", "com.tw", "co.il", "co.id", "com.my", "com.ph", "com.vn",
    "co.th", "com.pl", "com.co", "com.pe", "com.uy", "com.ua", "co.ke", "com.ng", "com.eg",
    "pages.dev", "github.io", "blogspot.com", "netlify.app", "vercel.app", "web.app",
    "firebaseapp.com", "herokuapp.com", "wordpress.com",
}


def _main_domain(domain: str) -> str:
    labels = domain.split(".")
    keep = 3 if ".".join(labels[-2:]) in MULTI_PART_SUFFIXES else 2
    return ".".join(labels[-keep:])


def _curl(url: str) -> dict:
    """Fetch a URL with curl like a browser would. Returns code, final URL and body."""
    cmd = ["curl", "-sS", "-L", "--max-redirs", "10", "--max-time", str(TIMEOUT_SECONDS),
           "--compressed", "-A", USER_AGENT, "-H", "Accept: text/plain,text/html;q=0.9,*/*;q=0.8",
           "-w", CURL_MARKER + "%{http_code} %{url_effective}", url]
    try:
        p = subprocess.run(cmd, capture_output=True, timeout=TIMEOUT_SECONDS + 5)
    except subprocess.TimeoutExpired:
        return {"code": None, "final_url": url, "body": "", "error": "timed out"}
    out = p.stdout.decode("utf-8", errors="replace")
    body, _, tail = out.rpartition(CURL_MARKER)
    code_text, _, final_url = tail.partition(" ")
    code = int(code_text) if code_text.isdigit() and code_text != "000" else None
    error = p.stderr.decode(errors="replace").strip().removeprefix("curl: ") if p.returncode else None
    return {"code": code, "final_url": final_url or url, "body": body, "error": error}


def _looks_like_html(body: str) -> bool:
    head = body[:3000].lower()
    return any(tag in head for tag in ("<!doctype", "<html", "<head", "<body"))


def _has_records(body: str) -> bool:
    """True if any non-comment line looks like an ads.txt record (fields separated by commas)."""
    return any("," in line.split("#", 1)[0] for line in body.splitlines())


def _classify_fetch(r: dict) -> str:
    """ads_txt, blocked, not_found or error. Sets r["note"] when the reason isn't obvious."""
    code, body = r["code"], r["body"]
    if code in (401, 403):
        return "blocked"
    if body and _looks_like_html(body) and any(m in body[:20000].lower() for m in BOT_CHECK_MARKERS):
        r["note"] = "bot check page"
        return "blocked"
    if code == 200:
        if _looks_like_html(body):
            r["note"] = "served an HTML page instead of ads.txt"
            return "not_found"
        if not _has_records(body):
            r["note"] = "ads.txt is empty or has no records"
            return "not_found"
        return "ads_txt"
    if code in (404, 410):
        return "not_found"
    return "error"


def _fetch_ads_txt(host: str) -> tuple[str, dict]:
    """Try https, then http. Keep the most useful outcome."""
    rank = {"ads_txt": 0, "blocked": 1, "not_found": 2, "error": 3}
    best = None
    for scheme in ("https", "http"):
        r = _curl(f"{scheme}://{host}/ads.txt")
        kind = _classify_fetch(r)
        if best is None or rank[kind] < rank[best[0]]:
            best = (kind, r)
        if kind == "ads_txt":
            break
    return best


def _match_onetag(body: str) -> tuple[str, str | None, str | None]:
    """Return (status, evidence line, relationship) for our line in an ads.txt body."""
    ours, ob, other = [], [], []
    for raw in body.splitlines():
        line = raw.split("#", 1)[0].strip()
        if not line:
            continue
        fields = [f.strip() for f in line.split(",")]
        if len(fields) < 2 or fields[0].lower() != ONETAG_DOMAIN:
            continue
        account = fields[1].lower()
        relationship = fields[2].upper() if len(fields) > 2 and fields[2] else None
        entry = (line, relationship)
        if account == ONETAG_ID:
            ours.append(entry)
        elif account == ONETAG_OB_ID:
            ob.append(entry)
        else:
            other.append(entry)
    for status, entries in (("authorised", ours), ("ob_only", ob), ("other_onetag_id", other)):
        if entries:
            direct = [e for e in entries if e[1] == "DIRECT"]
            line, relationship = (direct or entries)[0]
            return status, line, relationship
    return "no_onetag_line", None, None


def _check_host(host: str) -> dict:
    kind, r = _fetch_ads_txt(host)
    if kind == "blocked":
        time.sleep(BLOCK_RETRY_WAIT)
        kind, r = _fetch_ads_txt(host)
    result = {"checked_domain": host, "http_code": r["code"], "final_url": r["final_url"],
              "evidence": None, "relationship": None, "detail": r["error"]}
    if kind == "ads_txt":
        result["status"], result["evidence"], result["relationship"] = _match_onetag(r["body"])
    elif kind == "blocked":
        result["status"] = "blocked"
        result["detail"] = r.get("note") or f"HTTP {r['code']}"
    elif kind == "not_found":
        result["status"] = "no_ads_txt"
        result["detail"] = r.get("note") or f"HTTP {r['code']}"
    else:
        result["status"] = "error"
        result["detail"] = result["detail"] or f"HTTP {r['code']}"
    return result


def _check_one(domain: str) -> dict:
    """Check a site, falling back to its main domain if a subdomain has no ads.txt."""
    if domain.endswith(MIRROR_SUFFIX):
        result = {"checked_domain": domain, "used_main_domain": 0, "status": "mirror",
                  "http_code": None, "final_url": None, "evidence": None, "relationship": None,
                  "detail": f"translate.goog mirror of {_mirror_origin(domain)}"}
    else:
        result = _check_host(domain)
        result["used_main_domain"] = 0
        main = _main_domain(domain)
        if result["status"] == "no_ads_txt" and main != domain:
            fallback = _check_host(main)
            fallback["used_main_domain"] = 1
            fallback["detail"] = f"no ads.txt on {domain}; checked {main}" + (
                f" ({fallback['detail']})" if fallback["detail"] else "")
            result = fallback
    # When this site's check finished, not when the batch was saved.
    result["checked_at"] = datetime.now(timezone.utc).isoformat()
    return result


def _save_results(checked: list[tuple[str, float | None, dict]]) -> None:
    now = datetime.now(timezone.utc).isoformat()
    conn = _connect()
    with conn:
        for domain, bids, r in checked:
            conn.execute(
                "INSERT INTO results (domain, checked_domain, used_main_domain, status, http_code,"
                " final_url, evidence, relationship, detail, bid_requests, checked_at)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
                " ON CONFLICT(domain) DO UPDATE SET checked_domain = excluded.checked_domain,"
                " used_main_domain = excluded.used_main_domain, status = excluded.status,"
                " http_code = excluded.http_code, final_url = excluded.final_url,"
                " evidence = excluded.evidence, relationship = excluded.relationship,"
                " detail = excluded.detail, bid_requests = excluded.bid_requests,"
                " checked_at = excluded.checked_at",
                (domain, r["checked_domain"], r["used_main_domain"], r["status"], r["http_code"],
                 r["final_url"], r["evidence"], r["relationship"], r["detail"], bids,
                 r.get("checked_at") or now),
            )
            if r["status"] == "blocked":
                conn.execute(
                    "INSERT INTO manual_check (domain, bid_requests, url, http_code, reason, added_at)"
                    " VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(domain) DO UPDATE SET"
                    " bid_requests = excluded.bid_requests, url = excluded.url,"
                    " http_code = excluded.http_code, reason = excluded.reason",
                    (domain, bids, r["final_url"] or f"https://{r['checked_domain']}/ads.txt",
                     r["http_code"], r["detail"], now),
                )
            else:
                conn.execute("DELETE FROM manual_check WHERE domain = ?", (domain,))
    conn.close()


@tool
def check_domains(top_n: int = 50, offset: int = 0) -> str:
    """Check ads.txt for the next domains in bid-request order (highest first).

    Takes top_n domains from the domains table starting at offset, checks them
    about 20 at a time, and writes each result to the results table (blocked
    sites also go to manual_check). Returns only a count by status.
    """
    conn = _connect()
    targets = conn.execute(
        "SELECT domain, bid_requests FROM domains ORDER BY bid_requests DESC, domain"
        " LIMIT ? OFFSET ?", (int(top_n), int(offset))).fetchall()
    conn.close()

    with ThreadPoolExecutor(max_workers=CONCURRENCY) as pool:
        outcomes = list(pool.map(lambda t: _check_one(t["domain"]), targets))
    checked = [(t["domain"], t["bid_requests"], r) for t, r in zip(targets, outcomes)]
    _save_results(checked)

    return json.dumps({
        "checked": len(checked),
        "offset": int(offset),
        "by_status": dict(Counter(r["status"] for _, _, r in checked).most_common()),
        "used_main_domain": sum(r["used_main_domain"] for _, _, r in checked),
        "sent_to_manual_check": sum(r["status"] == "blocked" for _, _, r in checked),
    })


@tool
def check_domain(domain: str) -> str:
    """Re-check ads.txt for one site and update its row in the results table.

    Returns the status, the final URL, the matching line and the relationship,
    never the whole ads.txt file.
    """
    domain = _normalize_domain(domain)
    conn = _connect()
    row = conn.execute("SELECT bid_requests FROM domains WHERE domain = ?", (domain,)).fetchone()
    conn.close()
    r = _check_one(domain)
    _save_results([(domain, row["bid_requests"] if row else None, r)])
    return json.dumps({"domain": domain, **r})


# ---------- reading results and writing the brief ----------

RESULTS_DIR = ROOT / "results"
STATUSES = ("authorised", "ob_only", "other_onetag_id", "no_onetag_line",
            "no_ads_txt", "blocked", "mirror", "error")


def brief_path() -> Path:
    """Where today's brief is saved."""
    return RESULTS_DIR / f"brief-{datetime.now():%Y-%m-%d}.md"


def _connect_readonly() -> sqlite3.Connection:
    """Read-only connection: any write through it fails at the SQLite level."""
    conn = sqlite3.connect(DB_PATH.as_uri() + "?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


@tool
def query_results(status: str = "all", limit: int = 20) -> str:
    """Return checked domains from the results table, highest bid requests first.

    status is one of authorised, ob_only, other_onetag_id, no_onetag_line,
    no_ads_txt, blocked, mirror, error, or "all". limit is capped at 200.
    Read-only.
    """
    status = (status or "all").strip().lower()
    if status != "all" and status not in STATUSES:
        return json.dumps({"error": f"unknown status {status!r}; use one of {list(STATUSES)} or 'all'"})
    limit = max(1, min(int(limit), 200))
    sql = ("SELECT r.domain, COALESCE(d.bid_requests, r.bid_requests) AS bid_requests, r.status,"
           " r.checked_domain, r.used_main_domain, r.final_url, r.evidence, r.relationship,"
           " r.detail, r.checked_at FROM results r LEFT JOIN domains d USING (domain)")
    params: tuple = ()
    if status != "all":
        sql += " WHERE r.status = ?"
        params = (status,)
    sql += " ORDER BY bid_requests DESC LIMIT ?"
    conn = _connect_readonly()
    try:
        rows = [dict(r) for r in conn.execute(sql, (*params, limit))]
    finally:
        conn.close()
    return json.dumps({"status": status, "returned": len(rows), "rows": rows})


@tool
def summarise() -> str:
    """Counts by status and the bid requests in each, straight from the database.

    Also reports how many imported domains have been checked so far, how many
    sites are waiting in manual_check, how many authorised sites pass through
    Freestar's central file on a.pub.network, the DIRECT/RESELLER split
    among authorised sites, and the ten most common other onetag.com account
    IDs with their counts. Read-only.
    """
    conn = _connect_readonly()
    try:
        by_status = [dict(r) for r in conn.execute(
            "SELECT status, COUNT(*) AS domains, COALESCE(SUM(bid_requests), 0) AS bid_requests"
            " FROM results GROUP BY status ORDER BY bid_requests DESC")]
        imported = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(bid_requests), 0) AS b FROM domains").fetchone()
        manual = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(bid_requests), 0) AS b FROM manual_check").fetchone()
        freestar = conn.execute(
            "SELECT COUNT(*) AS n, COALESCE(SUM(bid_requests), 0) AS b FROM results"
            " WHERE status = 'authorised' AND final_url LIKE 'https://a.pub.network%'").fetchone()
        relationships = [dict(r) for r in conn.execute(
            "SELECT COALESCE(relationship, 'none') AS relationship, COUNT(*) AS domains,"
            " COALESCE(SUM(bid_requests), 0) AS bid_requests FROM results"
            " WHERE status = 'authorised' GROUP BY 1 ORDER BY bid_requests DESC")]
        other_ids = Counter()
        other_bids = Counter()
        for r in conn.execute("SELECT evidence, bid_requests FROM results"
                              " WHERE status = 'other_onetag_id' AND evidence IS NOT NULL"):
            fields = [f.strip() for f in r["evidence"].split(",")]
            account = fields[1].lower() if len(fields) > 1 else "?"
            other_ids[account] += 1
            other_bids[account] += r["bid_requests"] or 0
    finally:
        conn.close()
    checked = sum(s["domains"] for s in by_status)
    checked_bids = sum(s["bid_requests"] for s in by_status)
    return json.dumps({
        "by_status": by_status,
        "checked_domains": checked,
        "checked_bid_requests": checked_bids,
        "imported_domains": imported["n"],
        "imported_bid_requests": imported["b"],
        "unchecked_domains": imported["n"] - checked,
        "manual_check": {"domains": manual["n"], "bid_requests": manual["b"]},
        "authorised_via_freestar_file": {"domains": freestar["n"], "bid_requests": freestar["b"],
                                         "host": "a.pub.network"},
        "authorised_by_relationship": relationships,
        "other_onetag_ids": [{"account_id": a, "domains": n, "bid_requests": other_bids[a]}
                             for a, n in other_ids.most_common(10)],
    })


@tool
def write_brief(markdown: str) -> str:
    """Save a markdown brief to results/brief-YYYY-MM-DD.md (today's date), replacing any earlier one from today."""
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    path = brief_path()
    path.write_text(markdown, encoding="utf-8")
    return json.dumps({"saved": str(path.relative_to(ROOT)), "characters": len(markdown)})
